import argparse
import re
import sys
from pathlib import Path

import librosa
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from evaluation.evaluate_pitch_models import align_to_reference, build_reference_grid, compute_metrics
from modules.crepe_module import extract_crepe_pitch
from modules.pyin_module import extract_pyin_pitch_details, fuse_pyin_crepe


MODEL_ORDER = ["crepe", "pure_pyin", "hybrid_pyin"]
METRIC_ORDER = ["MAE_cents", "Raw_Pitch_Accuracy", "Voicing_Recall", "Voicing_False_Alarm"]
PYIN_FMIN = librosa.note_to_hz("C0")
PYIN_FMAX = librosa.note_to_hz("B8")
SHARP_WORDS = {
    "Csharp": "C#",
    "Dsharp": "D#",
    "Fsharp": "F#",
}

NOTE_TO_HZ = {
    "C0": 16.35, "C1": 32.70, "C2": 65.41, "C3": 130.81, "C4": 261.63, "C5": 523.25,
    "C6": 1046.50, "C7": 2093.00, "C8": 4186.01,
    "C#1": 34.65, "C#2": 69.30, "C#3": 138.59, "C#4": 277.18, "C#5": 554.37,
    "C#6": 1108.73, "C#7": 2217.46, "C#8": 4434.92,
    "E3": 164.81, "E4": 329.63, "E5": 659.25, "E6": 1318.51, "E7": 2637.02,
    "F#1": 46.25, "F#2": 92.50, "F#3": 185.00, "F#4": 369.99, "F#5": 739.99,
    "F#6": 1479.98, "F#7": 2959.96, "F#8": 5919.91,
    "G3": 196.00, "G4": 392.00, "G5": 783.99, "G6": 1567.98, "G7": 3135.96,
    "A3": 220.00, "A4": 440.00, "A5": 880.00, "A6": 1760.00, "A7": 3520.00,
    "B1": 61.74, "B2": 123.47, "B3": 246.94, "B4": 493.88, "B5": 987.77,
    "B6": 1975.53, "B7": 3951.07, "B8": 7902.13,
    "D#3": 155.56, "D#4": 311.13, "D#5": 622.25, "D#6": 1244.51, "D#7": 2489.02,
}
for _word, _symbol in SHARP_WORDS.items():
    for _octave in range(0, 9):
        _symbol_note = f"{_symbol}{_octave}"
        if _symbol_note in NOTE_TO_HZ:
            NOTE_TO_HZ[f"{_word}{_octave}"] = NOTE_TO_HZ[_symbol_note]


def infer_dataset(path: Path, dataset_root: Path) -> tuple[str, str]:
    rel_parts = path.relative_to(dataset_root).parts
    family = rel_parts[0].replace("_", " ").title() if rel_parts else "Unknown"
    subset = rel_parts[1] if len(rel_parts) > 2 else path.parent.name
    return family, subset


def normalize_note(note: str) -> str:
    note = note.replace("♯", "#").replace("＃", "#").strip()
    return note[0].upper() + note[1:]


def display_sharp_names(value: str) -> str:
    """Render Kaggle-safe sharp spellings as musical sharp symbols."""
    text = str(value)
    for word, symbol in SHARP_WORDS.items():
        text = re.sub(word, symbol, text, flags=re.IGNORECASE)
    return text


def parse_reference(source: str) -> tuple[float | None, str | None]:
    source = display_sharp_names(source)
    hz_match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*hz", source, flags=re.IGNORECASE)
    note_match = re.search(r"(?<![A-Za-z0-9])([A-Ga-g](?:#|b)?[0-8])(?![A-Za-z0-9])", source)

    note = normalize_note(note_match.group(1)) if note_match else None
    if hz_match:
        return float(hz_match.group(1)), note

    if note:
        if note in NOTE_TO_HZ:
            return float(NOTE_TO_HZ[note]), note
        try:
            return float(librosa.note_to_hz(note)), note
        except Exception:
            return None, note

    return None, None


def evaluate_audio_file(
    audio_path: Path,
    dataset_root: Path,
    crepe_conf_threshold: float,
    crepe_model_capacity: str,
) -> list[dict[str, float | str]]:
    ref_hz, ref_note = parse_reference(str(audio_path))
    if ref_hz is None:
        raise ValueError(f"Could not infer reference pitch from filename: {audio_path.name}")

    family, subset = infer_dataset(audio_path, dataset_root)
    crepe_raw_df = extract_crepe_pitch(
        str(audio_path),
        confidence_threshold=None,
        model_capacity=crepe_model_capacity,
    )
    crepe_df = crepe_raw_df.copy()
    crepe_df["frequency"] = np.where(
        crepe_df["confidence"].to_numpy(dtype=float) >= crepe_conf_threshold,
        crepe_df["frequency"].to_numpy(dtype=float),
        np.nan,
    )

    pyin_details_df = extract_pyin_pitch_details(str(audio_path), fmin=PYIN_FMIN, fmax=PYIN_FMAX)
    pyin_df = pyin_details_df.rename(columns={"pyin_f0": "hybrid_f0"})[["time", "hybrid_f0"]]
    hybrid_df = fuse_pyin_crepe(pyin_details_df, crepe_raw_df, fmin=PYIN_FMIN, fmax=PYIN_FMAX)
    ref_df = build_reference_grid(audio_path)

    model_frames = {
        "crepe": (crepe_df.rename(columns={"frequency": "pred_f0", "confidence": "pred_conf"}), "pred_conf"),
        "pure_pyin": (pyin_df.rename(columns={"hybrid_f0": "pred_f0"}), None),
        "hybrid_pyin": (hybrid_df.rename(columns={"hybrid_f0": "pred_f0"}), None),
    }

    rows = []
    for model in MODEL_ORDER:
        pred_df, conf_col = model_frames[model]
        aligned = align_to_reference(ref_df, pred_df, "pred_f0", conf_col)
        metrics = compute_metrics(aligned, model, ref_hz, crepe_conf_threshold)
        metrics.update({
            "Dataset": family,
            "Subset": subset,
            "File": audio_path.name,
            "Reference_Note": ref_note or "",
            "Reference_Hz": ref_hz,
            "Path": str(audio_path.relative_to(PROJECT_ROOT)),
        })
        rows.append(metrics)
    return rows


def summarize(results_df: pd.DataFrame, metric: str, agg: str) -> pd.DataFrame:
    table = results_df.pivot_table(
        index=["Dataset", "Subset"],
        columns="Model",
        values=metric,
        aggfunc=agg,
    ).reset_index()
    return table[["Dataset", "Subset", *[m for m in MODEL_ORDER if m in table.columns]]]


def best_per_note(results_df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for keys, group in results_df.groupby(["Dataset", "Subset", "File", "Reference_Note", "Reference_Hz"], dropna=False):
        valid = group.dropna(subset=["MAE_cents"])
        if valid.empty:
            best_model = ""
            best_mae = np.nan
        else:
            best = valid.loc[valid["MAE_cents"].idxmin()]
            best_model = best["Model"]
            best_mae = best["MAE_cents"]
        rows.append({
            "Dataset": keys[0],
            "Subset": keys[1],
            "File": keys[2],
            "Reference_Note": keys[3],
            "Reference_Hz": keys[4],
            "Best_Model": best_model,
            "Best_MAE_cents": best_mae,
        })
    return pd.DataFrame(rows).sort_values(["Dataset", "Subset", "Reference_Hz", "File"])


def safe_sheet_name(name: str, used: set[str]) -> str:
    name = display_sharp_names(name)
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Sheet"
    candidate = cleaned
    i = 2
    while candidate in used:
        suffix = f" {i}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
        i += 1
    used.add(candidate)
    return candidate


def add_old_vs_new_sheet(writer: pd.ExcelWriter, output_path: Path, old_path: Path | None) -> None:
    if old_path is None or not old_path.exists():
        pd.DataFrame([{"Status": "Old results workbook not found", "Expected_File": "Alpha_Testing_3_10_26__4_.xlsx"}]).to_excel(
            writer,
            sheet_name="Old vs New Hybrid",
            index=False,
        )
        return

    old_sheets = pd.read_excel(old_path, sheet_name=None)
    old_rows = []
    for sheet_name, sheet_df in old_sheets.items():
        sheet_df = sheet_df.copy()
        sheet_df.insert(0, "Old_Workbook_Sheet", sheet_name)
        old_rows.append(sheet_df)
    old_df = pd.concat(old_rows, ignore_index=True) if old_rows else pd.DataFrame()
    old_df.to_excel(writer, sheet_name="Old vs New Hybrid", index=False)


def write_workbook(results_df: pd.DataFrame, output_path: Path, old_path: Path | None) -> None:
    excel_df = results_df.copy()
    for col in ["Dataset", "Subset", "File", "Reference_Note"]:
        if col in excel_df.columns:
            excel_df[col] = excel_df[col].map(display_sharp_names)

    used_sheet_names: set[str] = set()
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summaries = {
            "Summary - Median MAE": summarize(excel_df, "MAE_cents", "median"),
            "Summary - Mean MAE": summarize(excel_df, "MAE_cents", "mean"),
            "Summary - Std Dev MAE": summarize(excel_df, "MAE_cents", "std"),
            "Summary - Raw Pitch Accuracy": summarize(excel_df, "Raw_Pitch_Accuracy", "mean"),
        }
        for sheet_name, df in summaries.items():
            used_sheet_names.add(sheet_name)
            df.to_excel(writer, sheet_name=sheet_name, index=False)

        for dataset, group in excel_df.groupby("Dataset", sort=True):
            sheet_name = safe_sheet_name(f"{dataset} Full Results", used_sheet_names)
            group.sort_values(["Subset", "Reference_Hz", "File", "Model"]).to_excel(writer, sheet_name=sheet_name, index=False)

        # Excel worksheet titles are limited to 31 characters.
        best_sheet_name = "Model Comparison - Best PerNote"
        best_per_note(excel_df).to_excel(writer, sheet_name=best_sheet_name, index=False)
        used_sheet_names.add(best_sheet_name)
        add_old_vs_new_sheet(writer, output_path, old_path)

    style_workbook(output_path)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    best_fill = PatternFill("solid", fgColor="C6EFCE")
    worst_fill = PatternFill("solid", fgColor="FFC7CE")

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        headers = [cell.value for cell in ws[1]]
        model_cols = [headers.index(model) + 1 for model in MODEL_ORDER if model in headers]
        lower_is_better = "Accuracy" not in ws.title
        if model_cols:
            for row in range(2, ws.max_row + 1):
                values = []
                for col in model_cols:
                    value = ws.cell(row=row, column=col).value
                    if isinstance(value, (int, float)) and np.isfinite(value):
                        values.append((col, value))
                if len(values) >= 2:
                    target_best = min(v for _, v in values) if lower_is_better else max(v for _, v in values)
                    target_worst = max(v for _, v in values) if lower_is_better else min(v for _, v in values)
                    for col, value in values:
                        if value == target_best:
                            ws.cell(row=row, column=col).fill = best_fill
                        if value == target_worst:
                            ws.cell(row=row, column=col).fill = worst_fill

        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value), 80))
            ws.column_dimensions[get_column_letter(col_idx)].width = max(10, max_len + 2)

    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Evaluate all sample_audio/datasets pitch datasets and export Excel summaries.")
    parser.add_argument("--dataset-root", default="sample_audio/datasets")
    parser.add_argument("--output", default="Full_Evaluation_Results.xlsx")
    parser.add_argument("--old-results", default="Alpha_Testing_3_10_26__4_.xlsx")
    parser.add_argument("--crepe-conf-threshold", type=float, default=0.5)
    parser.add_argument("--crepe-model-capacity", default="full", choices=["tiny", "small", "medium", "large", "full"])
    parser.add_argument("--checkpoint", default="Full_Evaluation_Checkpoint.csv")
    parser.add_argument("--finalize-only", action="store_true")
    parser.add_argument("--max-files", type=int, default=None)
    args = parser.parse_args()

    dataset_root = Path(args.dataset_root).resolve()
    if not dataset_root.exists():
        raise FileNotFoundError(f"Dataset root not found: {dataset_root}")

    wav_files = sorted(dataset_root.rglob("*.wav"))
    if args.max_files is not None:
        wav_files = wav_files[:args.max_files]
    if not wav_files:
        raise FileNotFoundError(f"No WAV files found under: {dataset_root}")

    checkpoint_path = PROJECT_ROOT / args.checkpoint
    rows = []
    completed_paths: set[str] = set()
    if checkpoint_path.exists():
        checkpoint_df = pd.read_csv(checkpoint_path)
        rows = checkpoint_df.to_dict("records")
        completed_paths = set(checkpoint_df["Path"].dropna().astype(str).unique()) if "Path" in checkpoint_df.columns else set()

    if args.finalize_only:
        if not rows:
            raise RuntimeError(f"No checkpoint rows found in {checkpoint_path}")
        write_workbook(pd.DataFrame(rows), PROJECT_ROOT / args.output, PROJECT_ROOT / args.old_results)
        print(f"Saved workbook to {PROJECT_ROOT / args.output}", flush=True)
        return

    skipped = []
    for idx, wav_path in enumerate(wav_files, start=1):
        rel_path = str(wav_path.relative_to(PROJECT_ROOT))
        if rel_path in completed_paths:
            print(f"[{idx}/{len(wav_files)}] Skipping completed {wav_path}", flush=True)
            continue

        print(f"[{idx}/{len(wav_files)}] Evaluating {wav_path}", flush=True)
        try:
            file_rows = evaluate_audio_file(
                wav_path,
                dataset_root,
                args.crepe_conf_threshold,
                args.crepe_model_capacity,
            )
            rows.extend(file_rows)
            pd.DataFrame(rows).to_csv(checkpoint_path, index=False)
            completed_paths.add(rel_path)
        except Exception as exc:
            skipped.append({"Path": str(wav_path), "Error": str(exc)})
            print(f"  skipped: {exc}", flush=True)

    if not rows:
        raise RuntimeError("No files were evaluated successfully.")

    results_df = pd.DataFrame(rows)
    output_path = PROJECT_ROOT / args.output
    old_path = PROJECT_ROOT / args.old_results
    write_workbook(results_df, output_path, old_path)

    if skipped:
        skipped_path = PROJECT_ROOT / "Full_Evaluation_Skipped.csv"
        pd.DataFrame(skipped).to_csv(skipped_path, index=False)
        print(f"Skipped {len(skipped)} files. Details saved to {skipped_path}", flush=True)

    print(f"Saved workbook to {output_path}", flush=True)


if __name__ == "__main__":
    main()
